"""Bounded document text extraction with immutable-version source locations.

Parsers never fetch HTML resources, evaluate spreadsheet formulas, call external
OCR/AI, or accept filenames as a retrieval instruction.
"""
from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from datetime import datetime, timezone
from html import unescape
from html.parser import HTMLParser
from typing import Protocol
from zipfile import BadZipFile, ZipFile

from .contracts import ParsedDocument, SourceDocument, TextBlock


PARSER_VERSION = "frip-documents-v1.1"


@dataclass(frozen=True)
class ParserLimits:
    max_bytes: int = 10 * 1024 * 1024
    max_characters: int = 1_000_000
    max_blocks: int = 20000
    max_pages: int = 200
    max_rows: int = 20000
    max_cells: int = 200000
    max_zip_entries: int = 1000
    max_uncompressed_bytes: int = 50 * 1024 * 1024
    max_compression_ratio: int = 200


@dataclass(frozen=True)
class OCRPage:
    page: int
    text: str
    confidence: float


class InternalOCRProvider(Protocol):
    """Implement only inside an approved corporate processing boundary."""
    version: str

    def extract(self, document: SourceDocument, max_pages: int) -> list[OCRPage]: ...


class _HTMLText(HTMLParser):
    BLOCK_TAGS = {"p","div","section","article","header","footer","li","ul","ol",
                  "h1","h2","h3","h4","h5","h6","table","tr","td","th","blockquote","pre","br","hr"}

    def __init__(self, source: str):
        super().__init__(convert_charrefs=False)
        self.source = source
        self.lines = source.splitlines(keepends=True)
        self.line_offsets = [0]
        for line in self.lines:
            self.line_offsets.append(self.line_offsets[-1] + len(line))
        self.ignored: list[str] = []
        self.pending: list[tuple[str, int, int]] = []
        self.parts: list[dict] = []

    def _flush(self):
        if not self.pending:
            return
        text = "".join(value for value,_,_ in self.pending)
        left = len(text) - len(text.lstrip())
        right = len(text.rstrip())
        if left < right:
            segments = []
            offset = 0
            for value,start,end in self.pending:
                overlap_start = max(left,offset)
                overlap_end = min(right,offset+len(value))
                if overlap_start < overlap_end:
                    if end-start == len(value):
                        raw_start = start+overlap_start-offset
                        raw_end = start+overlap_end-offset
                    else:
                        raw_start,raw_end = start,end
                    segments.append({"start":overlap_start-left,"end":overlap_end-left,
                                     "source_start":raw_start,"source_end":raw_end})
                offset += len(value)
            self.parts.append({"text":text[left:right],"source_start":segments[0]["source_start"],
                "source_end":segments[-1]["source_end"],"source_segments":segments,
                "locator":f"html:span:{segments[0]['source_start']}:{segments[-1]['source_end']}",
                "text_extraction_method":"html_text"})
        self.pending = []

    def handle_starttag(self, tag, attrs):
        if not self.ignored and tag in self.BLOCK_TAGS:
            self._flush()
        if tag in {"script", "style", "template", "iframe", "object", "noscript"}:
            self.ignored.append(tag)

    def handle_endtag(self, tag):
        if self.ignored and tag == self.ignored[-1]:
            self.ignored.pop()
        elif not self.ignored and tag in self.BLOCK_TAGS:
            self._flush()

    def _fragment(self, raw: str, rendered: str):
        if self.ignored:
            return
        line, offset = self.getpos()
        start = self.line_offsets[line - 1] + offset
        self.pending.append((rendered, start, start + len(raw)))

    def handle_data(self, data):
        self._fragment(data, data)

    def handle_entityref(self, name):
        raw = f"&{name};"
        self._fragment(raw, unescape(raw))

    def handle_charref(self, name):
        raw = f"&#{name};"
        self._fragment(raw, unescape(raw))

    def close(self):
        super().close()
        self._flush()


def _decode(content: bytes) -> str:
    return content.decode("utf-8-sig", errors="strict")


def _join_blocks(parts: list[dict], limits: ParserLimits) -> tuple[str, list[TextBlock]]:
    text_parts = []
    blocks = []
    offset = 0
    for index, item in enumerate(parts):
        value = item.pop("text")
        if offset + len(value) > limits.max_characters or len(blocks) >= limits.max_blocks:
            raise ValueError("document exceeds extracted text/block limits")
        if "source_segments" in item:
            item["source_segments"] = [segment | {"start":segment["start"]+offset,"end":segment["end"]+offset}
                                       for segment in item["source_segments"]]
        blocks.append(TextBlock(id=f"block-{index + 1}", text=value, start=offset, end=offset + len(value), **item))
        text_parts.append(value)
        offset += len(value) + 1
    return "\n".join(text_parts), blocks


def _text_blocks(text: str, limits: ParserLimits) -> list[TextBlock]:
    if len(text) > limits.max_characters:
        raise ValueError("document exceeds extracted text limits")
    blocks = []
    offset = 0
    for row, line in enumerate(text.splitlines(keepends=True), start=1):
        end = offset + len(line.rstrip("\r\n"))
        if text[offset:end].strip():
            blocks.append(TextBlock(id=f"block-{len(blocks)+1}", text=text[offset:end], start=offset, end=end,
                                    row=row, source_start=offset, source_end=end, locator=f"text:line:{row}"))
        offset += len(line)
        if len(blocks) > limits.max_blocks or row > limits.max_rows:
            raise ValueError("document exceeds line/block limits")
    return blocks


def _json_blocks(text: str, limits: ParserLimits) -> list[TextBlock]:
    # The standard decoder validates syntax; this walk retains exact scalar spans.
    decoder = json.JSONDecoder(parse_constant=lambda value: (_ for _ in ()).throw(ValueError("nonfinite JSON")))
    json.loads(text, parse_constant=lambda value: (_ for _ in ()).throw(ValueError("nonfinite JSON")))
    blocks = []

    def whitespace(pos):
        while pos < len(text) and text[pos].isspace():
            pos += 1
        return pos

    def walk(pos, path, depth=0):
        if depth > 64:
            raise ValueError("JSON nesting exceeds parser limit")
        pos = whitespace(pos)
        if text[pos] == "{":
            pos = whitespace(pos + 1)
            if text[pos] == "}":
                return pos + 1
            while True:
                key, end = decoder.raw_decode(text, pos)
                pos = whitespace(end)
                if text[pos] != ":":
                    raise ValueError("invalid JSON member")
                pos = walk(pos + 1, f"{path}.{key}", depth + 1)
                pos = whitespace(pos)
                if text[pos] == "}":
                    return pos + 1
                pos = whitespace(pos + 1)
        if text[pos] == "[":
            pos = whitespace(pos + 1)
            index = 0
            if text[pos] == "]":
                return pos + 1
            while True:
                pos = walk(pos, f"{path}[{index}]", depth + 1)
                index += 1
                pos = whitespace(pos)
                if text[pos] == "]":
                    return pos + 1
                pos = whitespace(pos + 1)
        _, end = decoder.raw_decode(text, pos)
        blocks.append(TextBlock(id=f"block-{len(blocks)+1}", text=text[pos:end], start=pos, end=end,
                                row=text.count("\n", 0, pos) + 1, source_start=pos, source_end=end, locator=f"json:{path}"))
        if len(blocks) > limits.max_blocks:
            raise ValueError("JSON exceeds scalar/block limits")
        return end

    walk(0, "$")
    return blocks


def _csv_blocks(text: str, limits: ParserLimits) -> list[TextBlock]:
    lines = text.splitlines(keepends=True)
    offsets = [0]
    for line in lines:
        offsets.append(offsets[-1] + len(line))
    reader = csv.reader(io.StringIO(text), strict=True)
    blocks = []
    previous_line = 0
    expected_columns = None
    for row_number, values in enumerate(reader, start=1):
        start = offsets[previous_line]
        end = offsets[reader.line_num]
        previous_line = reader.line_num
        raw = text[start:end].rstrip("\r\n")
        if row_number == 1:
            if not values or any(not value for value in values) or len(set(values)) != len(values):
                raise ValueError("CSV requires unique nonempty column names")
            expected_columns = len(values)
        elif len(values) != expected_columns:
            raise ValueError(f"CSV row {row_number} has inconsistent column count")
        blocks.append(TextBlock(id=f"block-{row_number}", text=raw, start=start, end=start+len(raw),
                                row=row_number, source_start=start, source_end=start+len(raw), locator=f"csv:row:{row_number}"))
        if row_number > limits.max_rows:
            raise ValueError("CSV exceeds row limits")
    return blocks


def _xlsx_parts(content: bytes, limits: ParserLimits) -> tuple[list[dict], list[str]]:
    with ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > limits.max_zip_entries:
            raise ValueError("spreadsheet archive has too many entries")
        total = sum(item.file_size for item in entries)
        compressed = sum(item.compress_size for item in entries)
        if total > limits.max_uncompressed_bytes or total > max(compressed, 1) * limits.max_compression_ratio:
            raise ValueError("spreadsheet archive exceeds decompression limits")
        if any(item.flag_bits & 1 for item in entries):
            raise ValueError("encrypted spreadsheet archives are unsupported")
        if any("vbaproject" in item.filename.casefold() for item in entries):
            raise ValueError("macro-enabled spreadsheet content is unsupported")
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    parts = []
    issues = []
    try:
        if sum((sheet.max_row or 0) * (sheet.max_column or 0) for sheet in book.worksheets) > limits.max_cells:
            raise ValueError("spreadsheet exceeds total scanned cell limits")
        for sheet in book.worksheets:
            if sheet.max_row and sheet.max_row > limits.max_rows:
                raise ValueError("spreadsheet exceeds row limits")
            if (sheet.max_row or 0) * (sheet.max_column or 0) > limits.max_cells:
                raise ValueError("spreadsheet exceeds scanned cell limits")
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    if cell.data_type == "f":
                        issues.append(f"formula not evaluated: {sheet.title}!{cell.coordinate}")
                        value = f"[FORMULA NOT EVALUATED] {cell.value}"
                    else:
                        value = str(cell.value)
                    parts.append({"text":value, "row":cell.row, "sheet":sheet.title,
                                  "locator":f"xlsx:{sheet.title}!{cell.coordinate}"})
                    if len(parts) > limits.max_blocks:
                        raise ValueError("spreadsheet exceeds cell/block limits")
    finally:
        book.close()
    return parts, issues


def parse_document(document: SourceDocument, *, limits: ParserLimits | None = None,
                   ocr_provider: InternalOCRProvider | None = None, now: datetime | None = None) -> ParsedDocument:
    limits = limits or ParserLimits()
    result = dict(source_system=document.source_system, original_id=document.original_id, revision=document.revision,
                  content_hash=document.content_hash, parser_version=PARSER_VERSION, source_timestamp=document.recorded_at,
                  effective_at=document.effective_at, ingested_at=now or datetime.now(timezone.utc),
                  equipment_alias=document.equipment_alias,chamber_alias=document.chamber_alias,collection=document.collection,
                  text="",blocks=[],status="parsed",issues=[])
    if len(document.content) > limits.max_bytes:
        return ParsedDocument(**(result | {"status":"error","issues":["document exceeds byte limit"]}))
    try:
        media = document.mime_type.split(";", 1)[0].strip().lower()
        if media in {"text/plain", "text/csv", "application/json"}:
            text = _decode(document.content)
            if len(text) > limits.max_characters:
                raise ValueError("document exceeds text limits")
            if media == "application/json":
                blocks = _json_blocks(text, limits)
            elif media == "text/csv":
                blocks = _csv_blocks(text, limits)
            else:
                blocks = _text_blocks(text, limits)
            result.update(text=text,blocks=blocks)
        elif media == "text/html":
            source = _decode(document.content)
            if len(source) > limits.max_characters:
                raise ValueError("HTML exceeds text limits")
            parser = _HTMLText(source)
            parser.feed(source)
            parser.close()
            result["text"], result["blocks"] = _join_blocks(parser.parts, limits)
        elif media == "application/pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(document.content), strict=True)
            if reader.is_encrypted:
                raise ValueError("encrypted PDF requires an approved internal decryption workflow")
            if len(reader.pages) > limits.max_pages:
                raise ValueError("PDF exceeds page limit")
            parts = []
            empty_pages = []
            for number,page in enumerate(reader.pages,start=1):
                text = page.extract_text() or ""
                if not text.strip():
                    empty_pages.append(number)
                offset = 0
                for row,line in enumerate(text.splitlines(keepends=True),start=1):
                    value = line.rstrip("\r\n")
                    if value.strip():
                        parts.append({"text":value,"page":number,"row":row,"source_start":offset,
                                      "source_end":offset+len(value),"locator":f"pdf:page:{number}:line:{row}"})
                    offset += len(line)
            if empty_pages and ocr_provider is None:
                result["status"] = "ocr_required"
                result["issues"].append(f"OCR required for pages {empty_pages}; no external OCR service is configured")
            elif empty_pages:
                pages = ocr_provider.extract(document, limits.max_pages)
                if len(pages) > limits.max_pages:
                    raise ValueError("OCR output exceeds page limit")
                if len({page.page for page in pages}) != len(pages):
                    raise ValueError("OCR output contains duplicate pages")
                for page in pages:
                    if page.page not in empty_pages or not 0 <= page.confidence <= 1:
                        raise ValueError("OCR output has invalid page or confidence")
                    parts.append({"text":page.text,"page":page.page,"locator":f"ocr:page:{page.page}",
                                  "text_extraction_method":"ocr","text_extraction_confidence":page.confidence,
                                  "text_extraction_provider_version":ocr_provider.version})
                result["issues"].append("Internal OCR output is unverified and requires human review")
                result["parser_version"] += "+ocr:" + ocr_provider.version
                if set(empty_pages) - {page.page for page in pages}:
                    result["status"] = "ocr_required"
            parts.sort(key=lambda part:(part.get("page",0),part.get("row",0)))
            result["text"], result["blocks"] = _join_blocks(parts, limits)
        elif media == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            parts,issues = _xlsx_parts(document.content, limits)
            result["text"], result["blocks"] = _join_blocks(parts, limits)
            result["issues"] += issues
        else:
            result.update(status="unsupported",issues=["unsupported document media type; install an approved parser adapter"])
    except (ValueError, TypeError, OSError, UnicodeError, BadZipFile, csv.Error, RecursionError) as exc:
        result.update(status="error",text="",blocks=[],issues=[f"document parser rejected input: {exc}"])
    except Exception as exc:
        # Third-party malformed-file errors vary by format; fail closed without content/log disclosure.
        result.update(status="error",text="",blocks=[],issues=[f"document parser rejected malformed input ({type(exc).__name__})"])
    return ParsedDocument(**result)
